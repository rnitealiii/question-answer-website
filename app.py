import os
import io
import csv
import sqlite3
import secrets
from functools import wraps
from difflib import SequenceMatcher

from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify
from werkzeug.utils import secure_filename

try:
    from rapidfuzz import fuzz, process
    HAS_RAPIDFUZZ = True
except ImportError:
    HAS_RAPIDFUZZ = False

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

APP_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get("QA_DATABASE_PATH", os.path.join(APP_DIR, "questions.db"))
ALLOWED_EXTENSIONS = {"csv", "xlsx"}
MATCH_THRESHOLD = int(os.environ.get("MATCH_THRESHOLD", "72"))

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY") or secrets.token_hex(32)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024  # 10 MB upload limit


def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS qa (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                question TEXT NOT NULL,
                answer TEXT NOT NULL,
                explanation TEXT NOT NULL DEFAULT '',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)


init_db()


def admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get("is_admin"):
            return redirect(url_for("admin_login", next=request.path))
        return fn(*args, **kwargs)
    return wrapper


def clean_text(value):
    if value is None:
        return ""
    return " ".join(str(value).strip().split())


def seed_from_bundled_csv():
    """Import the bundled GoTranscript Q&A CSV on first run (empty database)."""
    seed_csv = os.path.join(APP_DIR, "static", "GoTranscript-QA-Database.csv")
    if not os.path.exists(seed_csv):
        return
    with db() as conn:
        count = conn.execute("SELECT COUNT(*) FROM qa").fetchone()[0]
        if count > 0:
            return
        with open(seed_csv, encoding="utf-8-sig") as fh:
            reader = csv.DictReader(fh)
            rows = [
                (clean_text(r.get("Question")), clean_text(r.get("Answer")),
                 clean_text(r.get("Explanation")))
                for r in reader
            ]
            rows = [(q, a, e) for q, a, e in rows if q and a]
        conn.executemany(
            "INSERT INTO qa (question, answer, explanation) VALUES (?, ?, ?)", rows
        )


seed_from_bundled_csv()


def normalise_header(value):
    return "".join(ch.lower() for ch in clean_text(value) if ch.isalnum())


def map_headers(headers):
    aliases = {
        "question": {"question", "questions", "q", "prompt", "query"},
        "answer": {"answer", "answers", "a", "correctanswer", "solution"},
        "explanation": {"explanation", "explanations", "reason", "details", "rationale"},
    }
    normalized = [normalise_header(h) for h in headers]
    result = {}
    for field, names in aliases.items():
        for idx, header in enumerate(normalized):
            if header in names:
                result[field] = idx
                break
    return result


def parse_upload(file_storage):
    filename = secure_filename(file_storage.filename or "")
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    rows = []
    if ext == "csv":
        raw = file_storage.read()
        # Accept common CSV encodings.
        text = None
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("Could not read the CSV encoding.")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        all_rows = list(reader)
    elif ext == "xlsx":
        if load_workbook is None:
            raise ValueError("Excel support is not installed. Install the packages in requirements.txt.")
        workbook = load_workbook(file_storage, read_only=True, data_only=True)
        sheet = workbook.active
        all_rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
    else:
        raise ValueError("Please upload a .csv or .xlsx file.")

    if len(all_rows) < 2:
        raise ValueError("The file must contain a header row and at least one data row.")
    headers = list(all_rows[0])
    columns = map_headers(headers)
    if "question" not in columns or "answer" not in columns:
        raise ValueError("Could not identify Question and Answer columns. Use headers: Question, Answer, Explanation.")
    for row in all_rows[1:]:
        vals = list(row)
        def get(field):
            idx = columns.get(field)
            return clean_text(vals[idx]) if idx is not None and idx < len(vals) else ""
        question, answer, explanation = get("question"), get("answer"), get("explanation")
        if question and answer:
            rows.append((question, answer, explanation))
    if not rows:
        raise ValueError("No valid rows found. Each row needs a Question and an Answer.")
    return rows


def similarity(a, b):
    if HAS_RAPIDFUZZ:
        return float(fuzz.WRatio(a, b))
    return SequenceMatcher(None, a.casefold(), b.casefold()).ratio() * 100


def find_best_match(user_question):
    with db() as conn:
        records = conn.execute("SELECT id, question, answer, explanation FROM qa").fetchall()
    if not records:
        return None, 0
    query = clean_text(user_question)
    if not query:
        return None, 0

    # First handle exact case-insensitive matches.
    for record in records:
        if clean_text(record["question"]).casefold() == query.casefold():
            return record, 100

    questions = [record["question"] for record in records]
    if HAS_RAPIDFUZZ:
        best = process.extractOne(query, questions, scorer=fuzz.WRatio)
        if best:
            matched_question, score, idx = best
            return (records[idx], float(score)) if score >= MATCH_THRESHOLD else (None, float(score))
    best_record, best_score = None, 0
    for record in records:
        score = similarity(query, record["question"])
        if score > best_score:
            best_record, best_score = record, score
    return (best_record, best_score) if best_score >= MATCH_THRESHOLD else (None, best_score)


@app.route("/", methods=["GET"])
def index():
    return render_template("index.html")


@app.route("/ask", methods=["POST"])
def ask():
    question = clean_text(request.form.get("question", ""))
    if not question:
        return jsonify({"ok": False, "message": "Please enter a question."}), 400
    match, score = find_best_match(question)
    if match is None:
        return jsonify({
            "ok": True,
            "found": False,
            "message": "No matching answer found. Try rephrasing your question or ask an administrator to add it.",
        })
    return jsonify({
        "ok": True,
        "found": True,
        "question": match["question"],
        "answer": match["answer"],
        "explanation": match["explanation"] or "No explanation has been added for this question yet.",
        "confidence": round(score),
    })


@app.route("/admin")
def admin_home():
    if not session.get("is_admin"):
        return redirect(url_for("admin_login"))
    with db() as conn:
        count = conn.execute("SELECT COUNT(*) FROM qa").fetchone()[0]
        recent = conn.execute("SELECT id, question, answer, explanation FROM qa ORDER BY id DESC LIMIT 100").fetchall()
    return render_template("admin.html", count=count, entries=recent)


@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if request.method == "POST":
        expected = os.environ.get("ADMIN_PASSWORD")
        if not expected:
            flash("Admin login is not configured. Set the ADMIN_PASSWORD environment variable, then restart the app.", "error")
        elif secrets.compare_digest(request.form.get("password", ""), expected):
            session.clear()
            session["is_admin"] = True
            session.permanent = True
            return redirect(url_for("admin_home"))
        else:
            flash("Incorrect password.", "error")
    return render_template("login.html")


@app.route("/admin/logout", methods=["POST"])
def admin_logout():
    session.clear()
    return redirect(url_for("index"))


@app.route("/admin/upload", methods=["POST"])
@admin_required
def admin_upload():
    uploaded = request.files.get("file")
    if not uploaded or not uploaded.filename:
        flash("Choose a CSV or Excel file first.", "error")
        return redirect(url_for("admin_home"))
    try:
        rows = parse_upload(uploaded)
        with db() as conn:
            conn.executemany(
                "INSERT INTO qa (question, answer, explanation) VALUES (?, ?, ?)",
                rows
            )
        flash(f"Imported {len(rows)} question(s). Existing entries were kept.", "success")
    except Exception as exc:
        flash(str(exc), "error")
    return redirect(url_for("admin_home"))


@app.route("/admin/add", methods=["POST"])
@admin_required
def admin_add():
    question = clean_text(request.form.get("question", ""))
    answer = clean_text(request.form.get("answer", ""))
    explanation = clean_text(request.form.get("explanation", ""))
    if not question or not answer:
        flash("Question and answer are required.", "error")
    else:
        with db() as conn:
            conn.execute(
                "INSERT INTO qa (question, answer, explanation) VALUES (?, ?, ?)",
                (question, answer, explanation)
            )
        flash("Question added successfully.", "success")
    return redirect(url_for("admin_home"))


@app.route("/admin/delete/<int:entry_id>", methods=["POST"])
@admin_required
def admin_delete(entry_id):
    with db() as conn:
        conn.execute("DELETE FROM qa WHERE id = ?", (entry_id,))
    flash("Entry deleted.", "success")
    return redirect(url_for("admin_home"))


@app.route("/health")
def health():
    return jsonify({"status": "ok"})


if __name__ == "__main__":
    # For local testing only. Use a production WSGI server when deploying publicly.
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", "5000")), debug=False)
